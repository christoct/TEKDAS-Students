from pathlib import Path

import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
DATA = BASE / "data_prepared" / "coffee_sales_prepared.csv"
OUT = BASE / "Laporan_Analisis_Coffee_Sales.docx"
ASSET = BASE / "data_prepared"
df = pd.read_csv(DATA)

NAVY = (25, 43, 67)
TEAL = (8, 127, 140)
PURPLE = (118, 86, 181)
GOLD = (217, 139, 57)
INK = (40, 51, 65)
MUTED = (105, 119, 135)
BG = (246, 248, 251)
WHITE = (255, 255, 255)


def font(size, bold=False):
    candidates = [
        Path("C:/Windows/Fonts/segoeui.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
    ]
    if bold:
        candidates.insert(0, Path("C:/Windows/Fonts/segoeuib.ttf"))
        candidates.append(Path("C:/Windows/Fonts/arialbd.ttf"))
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def draw_dashboard(single_store):
    selected = df[df.store_location.eq("Astoria")].copy() if single_store else df.copy()
    width, height = 1500, 1040
    image = Image.new("RGB", (width, height), BG)
    d = ImageDraw.Draw(image)
    f_title, f_sub, f_body, f_small, f_metric = (
        font(33, True), font(17), font(17, True), font(14), font(29, True)
    )
    d.rectangle((0, 0, width, 95), fill=NAVY)
    d.text((42, 20), "Coffee Shop Sales", font=f_title, fill=WHITE)
    d.text((43, 62), "Dashboard • Januari–Juni 2023", font=f_sub, fill=(210, 222, 235))

    filter_text = "Filter gerai: Astoria" if single_store else "Filter gerai: Semua gerai"
    d.rounded_rectangle((42, 116, 405, 164), radius=12, fill=WHITE, outline=(222, 228, 235))
    d.text((60, 129), filter_text, font=f_body, fill=INK)
    d.text((width - 340, 130), "Periode: 2023-01 s.d. 2023-06", font=f_small, fill=MUTED)

    total = selected.sales.sum()
    transactions = selected.transaction_id.nunique()
    monthly = selected.groupby("month").agg(sales=("sales", "sum"), tx=("transaction_id", "nunique"))
    growth = monthly.sales.pct_change().iloc[-1] * 100
    avg = total / transactions
    metrics = [("Total penjualan", f"${total:,.2f}"), ("Transaksi tercatat", f"{transactions:,}"),
               ("Rata-rata nilai transaksi", f"${avg:,.2f}"), ("Pertumbuhan Jun vs Mei", f"{growth:.2f}%")]
    left, gap, card_w = 42, 18, 342
    for i, (label, value) in enumerate(metrics):
        x = left + i * (card_w + gap)
        d.rounded_rectangle((x, 186, x + card_w, 305), radius=16, fill=WHITE, outline=(222, 228, 235))
        d.text((x + 20, 207), label, font=f_small, fill=MUTED)
        d.text((x + 20, 245), value, font=f_metric, fill=NAVY)

    def panel(box, title):
        d.rounded_rectangle(box, radius=16, fill=WHITE, outline=(222, 228, 235))
        d.text((box[0] + 22, box[1] + 16), title, font=f_body, fill=NAVY)

    panel((42, 327, 740, 650), "Tren penjualan bulanan")
    panel((760, 327, 1458, 650), "Penjualan per gerai")
    plot_x0, plot_y0, plot_x1, plot_y1 = 92, 390, 690, 595
    d.line((plot_x0, plot_y1, plot_x1, plot_y1), fill=(190, 200, 210), width=2)
    vals = monthly.sales.tolist()
    max_val = max(vals) * 1.12
    points = []
    for i, val in enumerate(vals):
        x = plot_x0 + i * ((plot_x1 - plot_x0) / 5)
        y = plot_y1 - (val / max_val) * (plot_y1 - plot_y0)
        points.append((x, y))
        d.text((x - 25, plot_y1 + 11), monthly.index[i][5:] if False else str(monthly.index[i])[-2:], font=f_small, fill=MUTED)
    d.line(points, fill=TEAL, width=5)
    for x, y in points:
        d.ellipse((x - 7, y - 7, x + 7, y + 7), fill=TEAL)

    stores = selected.groupby("store_location").sales.sum().sort_values(ascending=False)
    max_store = stores.max() or 1
    for i, (name, value) in enumerate(stores.items()):
        y = 396 + i * 62
        d.text((790, y), name, font=f_small, fill=INK)
        bar_x = 960
        bar_w = int(380 * value / max_store)
        d.rounded_rectangle((bar_x, y, bar_x + bar_w, y + 26), radius=8,
                            fill=[TEAL, PURPLE, GOLD][i % 3])
        d.text((bar_x + bar_w + 12, y + 3), f"${value:,.0f}", font=f_small, fill=INK)

    panel((42, 672, 740, 1007), "Penjualan per kategori")
    panel((760, 672, 1458, 1007), "Transaksi tercatat menurut jam")
    cat = selected.groupby("product_category").sales.sum().sort_values(ascending=False)
    colors = [TEAL, PURPLE, GOLD, (73, 156, 111), (200, 99, 91), (119, 143, 168), (225, 190, 89)]
    cx, cy, radius, thickness = 217, 837, 112, 48
    start = -90
    for i, value in enumerate(cat.values):
        extent = value / cat.sum() * 360
        d.arc((cx-radius, cy-radius, cx+radius, cy+radius), start, start + extent,
              fill=colors[i % len(colors)], width=thickness)
        start += extent
    d.text((cx - 52, cy - 13), "Kategori", font=f_small, fill=MUTED)
    for i, (name, value) in enumerate(cat.head(5).items()):
        yy = 738 + i * 48
        d.ellipse((385, yy + 5, 401, yy + 21), fill=colors[i % len(colors)])
        d.text((414, yy), name, font=f_small, fill=INK)
        d.text((414, yy + 20), f"${value:,.2f} • {value/cat.sum():.1%}", font=f_small, fill=MUTED)

    hours = selected.groupby("hour").transaction_id.nunique().sort_values(ascending=False).head(10).sort_index()
    x0, y_base = 820, 925
    max_hour = hours.max()
    for i, (hour, count) in enumerate(hours.items()):
        x = x0 + i * 59
        bar_h = int(145 * count / max_hour)
        d.rectangle((x, y_base - bar_h, x + 33, y_base), fill=GOLD if count == max_hour else TEAL)
        d.text((x - 1, y_base + 10), f"{int(hour):02d}", font=f_small, fill=MUTED)
    d.text((802, 970), f"Puncak: {int(selected.groupby('hour').transaction_id.nunique().idxmax()):02d}:00", font=f_small, fill=INK)
    d.text((44, 1018), "Nilai mata uang tidak dinyatakan pada sumber; $ hanya penanda format angka sumber.", font=f_small, fill=MUTED)
    path = ASSET / ("dashboard_astoria.png" if single_store else "dashboard_semua_gerai.png")
    image.save(path, quality=95)
    return path


astoria = df[df.store_location.eq("Astoria")]
store_sales = df.groupby("store_location").sales.sum().sort_values(ascending=False)
cats = df.groupby("product_category").sales.sum().sort_values(ascending=False)
products = df.groupby(["product_id", "product_detail"]).sales.sum().sort_values(ascending=False)
hour_counts = df.groupby("hour").transaction_id.nunique()
astoria_month = astoria.groupby("month").agg(sales=("sales", "sum"), tx=("transaction_id", "nunique"))
images = [draw_dashboard(True), draw_dashboard(False)]


def rupiah_like(value):
    return f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)
styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10)
styles["Normal"].font.color.rgb = RGBColor(*INK)
for stylename in ("Heading 1", "Heading 2"):
    styles[stylename].font.name = "Aptos Display"
    styles[stylename].font.color.rgb = RGBColor(*NAVY)

title = doc.add_heading("Analisis Dashboard Penjualan Coffee Shop", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph("Periode 1 Januari–30 Juni 2023 | Sumber: coffee_sales_prepared.csv dan monthly_summary.csv")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph("Catatan: berkas sumber tidak menyatakan mata uang. Angka penjualan di bawah ditulis sebagai nilai numerik dengan pemisah ribuan desimal Indonesia.")

doc.add_heading("1. Definisi empat indikator", level=1)
table = doc.add_table(rows=1, cols=3)
table.style = "Light Shading Accent 1"
for cell, value in zip(table.rows[0].cells, ["Indikator", "Definisi", "Kolom / rumus"]):
    cell.text = value
definitions = [
    ("Total penjualan", "Jumlah nilai penjualan pada baris transaksi dalam cakupan periode dan gerai terpilih.", "sales dijumlahkan (total_sales)"),
    ("Transaksi tercatat", "Jumlah ID transaksi unik; pada data ini satu baris mewakili satu transaction_id dan satu produk.", "nunique(transaction_id) (total_transactions)"),
    ("Rata-rata nilai transaksi", "Total penjualan dibagi jumlah transaction_id unik. Angka ini belum dipastikan sebagai nilai seluruh isi struk pelanggan.", "total_sales ÷ total_transactions (average_transaction_value)"),
    ("Pertumbuhan penjualan bulanan", "Persentase perubahan total penjualan bulan terakhir dibanding bulan sebelumnya.", "(bulan ini − bulan lalu) ÷ bulan lalu × 100% (growth_pct)"),
]
for row in definitions:
    cells = table.add_row().cells
    for cell, value in zip(cells, row):
        cell.text = value

doc.add_heading("2. Filter Astoria", level=1)
may, june = astoria_month.loc["2023-05"], astoria_month.loc["2023-06"]
doc.add_paragraph(f"Dengan hanya Astoria dipilih, total penjualan selama Jan–Jun adalah {rupiah_like(astoria.sales.sum())} dan terdapat {astoria.transaction_id.nunique():,} transaksi tercatat. Untuk mencatat perubahan pada cakupan yang sama: Mei {rupiah_like(may.sales)} dengan {may.tx:,} transaksi; Juni {rupiah_like(june.sales)} dengan {june.tx:,} transaksi. Dari Mei ke Juni, penjualan naik {((june.sales/may.sales)-1)*100:.2f}% dan transaksi naik {((june.tx/may.tx)-1)*100:.2f}%.")

doc.add_heading("3. Perbandingan tiga gerai", level=1)
rank = "; ".join(f"{name}: {rupiah_like(value)}" for name, value in store_sales.items())
doc.add_paragraph(f"Setelah semua gerai dipilih kembali, grafik batang menunjukkan Hell's Kitchen tertinggi ({rupiah_like(store_sales.iloc[0])}), diikuti Astoria ({rupiah_like(store_sales.iloc[1])}) dan Lower Manhattan ({rupiah_like(store_sales.iloc[2])}). Selisih Hell's Kitchen terhadap Astoria sekitar {((store_sales.iloc[0]/store_sales.iloc[1])-1)*100:.2f}%. Nilai adalah akumulasi periode Jan–Jun.")

doc.add_heading("4. Kategori terbesar", level=1)
top_category = cats.index[0]
doc.add_paragraph(f"Pada cakupan semua gerai dan seluruh periode, kategori terbesar adalah {top_category}: {rupiah_like(cats.iloc[0])}, setara {cats.iloc[0]/df.sales.sum()*100:.2f}% dari total penjualan {rupiah_like(df.sales.sum())}. Karena kategori memakai cakupan semua gerai, angka dan persentase harus dibaca dengan filter semua gerai.")

doc.add_heading("5. Produk teratas", level=1)
top_product = products.index[0]
doc.add_paragraph(f"Contoh produk teratas: {top_product[1]} (ID {top_product[0]}), dengan total penjualan {rupiah_like(products.iloc[0])} di seluruh gerai selama Jan–Jun. Grafik mengurutkan produk berdasarkan nilai penjualan (sales), bukan jumlah unit terjual ataupun keuntungan.")

doc.add_heading("6. Jam transaksi tertinggi", level=1)
peak_hour, peak_count = hour_counts.idxmax(), hour_counts.max()
doc.add_paragraph(f"Puncak transaksi tercatat terjadi pada pukul {int(peak_hour):02d}:00, sebanyak {peak_count:,} transaction_id unik sepanjang periode. Ini belum membuktikan antrean panjang: grafik hanya menghitung transaksi menurut jam, tanpa data waktu tunggu, panjang antrean, durasi pelayanan, atau jumlah pelanggan yang sedang menunggu. Perbedaan jumlah staf dan laju kedatangan juga tidak tercakup.")

doc.add_heading("7. Insight dan tindak lanjut", level=1)
doc.add_paragraph(f"Pada periode Mei ke Juni 2023 dengan filter hanya Astoria, penjualan meningkat dari {rupiah_like(may.sales)} menjadi {rupiah_like(june.sales)} (+{((june.sales/may.sales)-1)*100:.2f}%), sementara transaksi meningkat dari {may.tx:,} menjadi {june.tx:,} (+{((june.tx/may.tx)-1)*100:.2f}%). Kenaikan keduanya menunjukkan aktivitas penjualan Astoria bertambah pada bulan tersebut, tetapi data tidak menjelaskan penyebabnya dan satu baris transaksi belum tentu satu struk pelanggan penuh. Tindakan pemeriksaan: bandingkan data harian dan produk terjual untuk Mei–Juni, lalu cocokkan dengan catatan promosi, stok, dan jadwal staf sebelum menetapkan tindakan operasional.")

doc.add_heading("8. Tangkapan layar dashboard", level=1)
doc.add_paragraph("Dua visual dashboard berikut disusun dari CSV untuk memperlihatkan cakupan filter yang berbeda; gambar ini bukan tangkapan layar langsung dari aplikasi. Nilai memakai data sumber, yang tidak menyatakan mata uang.")
for caption, path in [("Visual 1 — filter hanya Astoria", images[0]), ("Visual 2 — semua gerai dipilih", images[1])]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(caption)
    run.bold = True
    run.font.color.rgb = RGBColor(*NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(6.95))

doc.add_heading("9. Laporan perubahan dan penambahan fitur", level=1)
doc.add_paragraph(
    "Dashboard diperluas dengan machine learning (ML) untuk membuat proyeksi penjualan dan "
    "large language model (LLM) untuk merangkum evidence menjadi opsi tindak lanjut. Fitur ini "
    "mendukung analisis manajemen; keputusan akhir tetap memerlukan pemeriksaan data operasional."
)
feature_table = doc.add_table(rows=1, cols=3)
feature_table.style = "Light Shading Accent 1"
for cell, value in zip(feature_table.rows[0].cells, ["Fitur", "Fungsi", "Kegunaan bagi owner / manajemen"]):
    cell.text = value
feature_rows = [
    (
        "Forecast ML — Random Forest",
        "Menggunakan kalender serta lag penjualan harian untuk memproyeksikan penjualan 7, 14, 21, atau 28 hari. Riwayat mengikuti gerai, kategori, hari, dan jam terpilih hingga tanggal akhir filter.",
        "Membantu menyusun skenario kebutuhan operasional dan membandingkan kemungkinan arah penjualan pada cakupan filter yang sama. Forecast dapat diunduh sebagai CSV.",
    ),
    (
        "Backtest dan baseline mingguan",
        "Menguji prediksi satu hari ke depan pada bagian akhir data kronologis. Menampilkan MAE dan RMSE Random Forest dibanding baseline yang mengulang penjualan pada hari yang sama minggu sebelumnya.",
        "Memberi pemeriksaan sederhana apakah model lebih akurat daripada acuan historis. Jika baseline lebih baik, dashboard memberi peringatan agar forecast ML tidak terlalu diandalkan.",
    ),
    (
        "Asisten analisis LLM via Ollama",
        "Mengirim ringkasan terfilter—KPI, tren bulanan, gerai, kategori, produk teratas, hari/jam sibuk, pembanding periode, dan forecast bila ada—ke model Ollama. Evidence dan prompt dapat diperiksa sebelum analisis dibuat.",
        "Meringkas temuan dalam bahasa yang mudah dibaca dan memberi opsi tindak lanjut beserta pemeriksaan data yang disarankan. LLM diarahkan agar membedakan fakta dari hipotesis dan tidak mengarang penyebab.",
    ),
    (
        "Kontrol penggunaan dan batas interpretasi",
        "Forecast dan analisis LLM dijalankan saat diminta. Rekomendasi memakai filter aktif; perubahan filter membuat hasil lama ditandai agar diperbarui. Model LLM dapat dipilih dari aplikasi.",
        "Mengurangi risiko membaca saran dari cakupan lama sebagai hasil terbaru. Rekomendasi tetap perlu divalidasi, bukan dijalankan sebagai keputusan otomatis.",
    ),
]
for row in feature_rows:
    cells = feature_table.add_row().cells
    for cell, value in zip(cells, row):
        cell.text = value

doc.add_heading("Cara memakai fitur baru", level=2)
doc.add_paragraph(
    "1. Buka tab Forecast ML, atur horizon, lalu pilih Latih dan buat forecast. Tinjau MAE/RMSE "
    "bersama baseline sebelum menggunakan hasil untuk perencanaan.\n"
    "2. Buka tab Analisis LLM dan periksa evidence pada expander. Setelah Ollama berjalan dan "
    "model tersedia, pilih Buat analisis keputusan. Tinjau rekomendasi serta data pemeriksaannya."
)
doc.add_paragraph(
    "Prasyarat LLM: pasang dependensi dari Pertemuan-5-BI/requirements.txt, jalankan layanan "
    "Ollama, dan unduh model default llama3.2 (atau masukkan model lain yang tersedia). ML "
    "memerlukan scikit-learn. Petunjuk setup tersedia di Pertemuan-5-BI/README.md."
)

doc.add_heading("Batasan ML dan LLM", level=2)
doc.add_paragraph(
    "Riwayat yang tersedia hanya Januari–Juni 2023, sehingga pola tahunan, hari libur, perubahan "
    "harga, promosi, stok, biaya, margin, dan status operasional belum dimodelkan. Backtest satu "
    "hari ke depan tidak menjamin akurasi forecast multi-hari. Sumber juga tidak menyatakan mata "
    "uang. LLM hanya membantu menafsirkan ringkasan; verifikasi penjualan, persediaan, margin, "
    "promosi, antrean, dan jadwal staf sebelum menetapkan tindakan bisnis."
)

doc.save(OUT)
print(OUT)
